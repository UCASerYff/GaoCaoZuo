#!/usr/bin/env python3
"""Generate original empty document templates with open-source document libraries.
Development only; the shipped application has no Python dependency.
"""
from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
import zipfile
import struct

root = Path(__file__).resolve().parents[1] / "Resources" / "DocumentTemplates"
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.core_properties.title = ""
doc.add_paragraph("")
doc.save(root / "Blank.docx")
book = Workbook()
book.properties.creator = ""
book.properties.lastModifiedBy = ""
book.active.title = "Sheet1"
book.save(root / "Blank.xlsx")
deck = Presentation()
deck.core_properties.author = ""
deck.core_properties.last_modified_by = ""
deck.slides.add_slide(deck.slide_layouts[6])
deck.save(root / "Blank.pptx")
(root / "Blank.rtf").write_text(r"{\rtf1\ansi\deff0{\fonttbl{\f0 Helvetica;}}\f0\fs24\par}", encoding="ascii")
(root / "Blank.xml").write_text('<?xml version="1.0" encoding="UTF-8"?>\n<document/>\n', encoding="utf-8")
# Original 512×512 RGB white canvas. Native PSD, not a renamed image.
# Adobe's public PSD specification: big-endian 26-byte header, three empty
# length-delimited sections, compression=0, then planar 8-bit RGB image data.
# https://www.adobe.com/devnet-apps/photoshop/fileformatashtml/
size = 512
psd = struct.pack(">4sH6sHIIHH", b"8BPS", 1, bytes(6), 3, size, size, 8, 3)
psd += struct.pack(">IIIH", 0, 0, 0, 0) + bytes([255]) * (size * size * 3)
(root / "Blank.psd").write_bytes(psd)
Document(root / "Blank.docx")
load_workbook(root / "Blank.xlsx")
Presentation(root / "Blank.pptx")
for file in root.glob("*.???x"):
    with zipfile.ZipFile(file) as archive:
        assert archive.testzip() is None
print("Generated and reopened DOCX/XLSX/PPTX; original RTF/XML/PSD templates ready.")
